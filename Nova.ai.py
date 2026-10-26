import os
import shutil
import subprocess
import tempfile
import sqlite3
import threading
import requests
import telebot
import io
import base64
import asyncio
import edge_tts
import time
import csv
import re
import json
import difflib
from datetime import datetime, timezone
from PyPDF2 import PdfReader
import pandas as pd
import pdf2docx
from PIL import Image
from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from telebot import types
from http.server import BaseHTTPRequestHandler, HTTPServer

# =========================================================
# НАСТРОЙКИ
# =========================================================

BOT_TOKEN = os.environ.get("BOT_TOKEN")
GROQ_KEY = os.environ.get("GROQ_KEY")
OPENROUTER_KEY = os.environ.get("OPENROUTER_KEY")
ADMIN_ID = int(os.environ.get("ADMIN_ID", "0"))
MODEL = os.environ.get("GROQ_MODEL", "openai/gpt-oss-120b")

# Настройки GitHub
GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN")
GITHUB_REPO = os.environ.get("GITHUB_REPO", "igorbobyrev72-sketch/Nova-ai")
GITHUB_BASE_BRANCH = os.environ.get("GITHUB_BASE_BRANCH", "main")
GITHUB_API = "https://api.github.com"
CODE_FILE = "Nova.ai.py"

if not BOT_TOKEN or not GROQ_KEY:
    raise SystemExit("Ошибка: не заданы BOT_TOKEN или GROQ_KEY")

bot = telebot.TeleBot(BOT_TOKEN)
user_files = {}
find_error_mode = {}
pending_code_changes = {}

# =========================================================
# БАЗА ДАННЫХ
# =========================================================

def init_db():
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS users (chat_id INTEGER PRIMARY KEY, role TEXT DEFAULT 'Обычный', model TEXT, requests INTEGER DEFAULT 0)''')
    c.execute('''CREATE TABLE IF NOT EXISTS history (id INTEGER PRIMARY KEY AUTOINCREMENT, chat_id INTEGER, role TEXT, content TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS favorites (id INTEGER PRIMARY KEY AUTOINCREMENT, chat_id INTEGER, text TEXT)''')
    c.execute('''CREATE TABLE IF NOT EXISTS core_tasks (id INTEGER PRIMARY KEY AUTOINCREMENT, chat_id INTEGER, task_text TEXT, agent_type TEXT, result TEXT, status TEXT DEFAULT 'pending')''')
    conn.commit()
    conn.close()

init_db()

def get_user(chat_id):
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("SELECT role, model, requests FROM users WHERE chat_id = ?", (chat_id,))
    row = c.fetchone()
    if not row:
        c.execute("INSERT INTO users (chat_id, model) VALUES (?, ?)", (chat_id, MODEL))
        conn.commit()
        row = ("Обычный", MODEL, 0)
    conn.close()
    return row

def update_user(chat_id, role=None, model=None):
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    if role: c.execute("UPDATE users SET role = ? WHERE chat_id = ?", (role, chat_id))
    if model: c.execute("UPDATE users SET model = ? WHERE chat_id = ?", (model, chat_id))
    conn.commit()
    conn.close()

def inc_requests(chat_id):
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("UPDATE users SET requests = requests + 1 WHERE chat_id = ?", (chat_id,))
    conn.commit()
    conn.close()

def get_history(chat_id):
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("SELECT role, content FROM history WHERE chat_id = ? ORDER BY id DESC LIMIT 10", (chat_id,))
    rows = c.fetchall()
    conn.close()
    return [{"role": r[0], "content": r[1]} for r in reversed(rows)]

def add_history(chat_id, role, content):
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("INSERT INTO history (chat_id, role, content) VALUES (?, ?, ?)", (chat_id, role, content))
    conn.commit()
    conn.close()

def clear_history(chat_id):
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("DELETE FROM history WHERE chat_id = ?", (chat_id,))
    conn.commit()
    conn.close()

# =========================================================
# HEALTH CHECK ДЛЯ RENDER
# =========================================================

class Health(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(b"ok")
    def log_message(self, *args): pass

def run_health():
    port = int(os.environ.get("PORT", "8080"))
    HTTPServer(("0.0.0.0", port), Health).serve_forever()

# =========================================================
# ОСНОВНАЯ ЛОГИКА
# =========================================================

def is_allowed(m):
    if ADMIN_ID != 0 and m.chat.id != ADMIN_ID:
        bot.reply_to(m, "Доступ закрыт.")
        return False
    return True

def get_main_keyboard(chat_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    role, _, _ = get_user(chat_id)
    err_mode = "ВКЛ" if find_error_mode.get(chat_id, False) else "ВЫКЛ"
    markup.add(types.KeyboardButton(f"Роль: {role}"), types.KeyboardButton(f"Найди ошибку: {err_mode}"))
    markup.add(types.KeyboardButton("Сменить модель"), types.KeyboardButton("Профиль"))
    markup.add(types.KeyboardButton("Очистить историю"), types.KeyboardButton("Поиск"))
    return markup

def get_main_inline_menu():
    markup = types.InlineKeyboardMarkup(row_width=2)
    markup.add(types.InlineKeyboardButton("💬 ИИ-Чат", callback_data="menu_chat"), types.InlineKeyboardButton("🧠 Режимы", callback_data="menu_modes"))
    markup.add(types.InlineKeyboardButton("📜 История", callback_data="menu_history"), types.InlineKeyboardButton("👤 Профиль", callback_data="menu_profile"))
    markup.add(types.InlineKeyboardButton("⚙️ Настройки", callback_data="menu_settings"), types.InlineKeyboardButton("❓ Помощь", callback_data="menu_help"))
    return markup

def send_smart_response(chat_id, text, reply_to_message_id=None):
    code_blocks = re.findall(r"```(?:python)?\s*\n(.*?)```", text, re.DOTALL)
    if code_blocks and len(code_blocks[0]) > 500:
        clean_text = re.sub(r"```(?:python)?\s*\n(.*?)```", "[Код отправлен файлом ниже]", text, flags=re.DOTALL)
        bot.send_message(chat_id, clean_text[:4000], reply_to_message_id=reply_to_message_id)
        doc = io.BytesIO("\n\n".join(code_blocks).encode("utf-8"))
        doc.name = "nova_code.py"
        bot.send_document(chat_id, doc, caption="📁 Код Nova SI")
    elif len(text) > 4000:
        doc = io.BytesIO(text.encode("utf-8"))
        doc.name = "nova_text.txt"
        bot.send_document(chat_id, doc, caption="📁 Текст слишком длинный для одного сообщения.")
    else:
        bot.send_message(chat_id, text, reply_to_message_id=reply_to_message_id)

def ask_groq(user_id, text, system_prompt=None):
    history = get_history(user_id)
    if len(history) > 6: history = history[-6:]
    history.append({"role": "user", "content": text[:3000]})
    role, _, _ = get_user(user_id)

    if system_prompt: sys_prompt = system_prompt
    elif role == "Дебаггер": sys_prompt = "Ты — сеньор-разработчик. Найди ошибку, объясни её и предложи исправление."
    elif role == "Эксперт": sys_prompt = "Ты — строгий эксперт. Отвечай точно и без воды."
    elif role == "Юморист": sys_prompt = "Ты — юморист. Отвечай с уместными шутками."
    elif role == "Программист": sys_prompt = "Ты — опытный программист. Пиши рабочий код."
    elif role == "Дружелюбный": sys_prompt = "Ты — дружелюбный и полезный помощник."
    else: sys_prompt = "Ты — Nova SI, полезный ИИ-помощник. Отвечай по-русски. Для кода используй Markdown-блоки."

    messages = [{"role": "system", "content": sys_prompt}] + history
    _, current_model, _ = get_user(user_id)

    url = "https://api.groq.com/openai/v1/chat/completions"
    headers = {"Authorization": f"Bearer {GROQ_KEY}", "Content-Type": "application/json"}
    data = {"model": current_model, "messages": messages}

    response = requests.post(url, headers=headers, json=data, timeout=60)
    response.raise_for_status()
    answer = response.json()["choices"][0]["message"]["content"]
    if not answer: raise ValueError("Модель вернула пустой ответ.")

    add_history(user_id, "user", text[:3000])
    add_history(user_id, "assistant", answer)
    inc_requests(user_id)
    return answer

# =========================================================
# NOVA CODE — GITHUB ИЗМЕНЕНИЯ
# =========================================================

def github_headers():
    if not GITHUB_TOKEN: raise RuntimeError("Не задан GITHUB_TOKEN в переменных Render.")
    return {"Authorization": f"Bearer {GITHUB_TOKEN}", "Accept": "application/vnd.github+json", "X-GitHub-Api-Version": "2022-11-28"}

def github_request(method, path, **kwargs):
    response = requests.request(method, f"{GITHUB_API}{path}", headers=github_headers(), timeout=30, **kwargs)
    if not response.ok: raise RuntimeError(f"GitHub API: HTTP {response.status_code}: {response.text[:500]}")
    if response.status_code == 204: return {}
    return response.json()

def github_get_file():
    path = f"/repos/{GITHUB_REPO}/contents/{CODE_FILE}"
    result = github_request("GET", path, params={"ref": GITHUB_BASE_BRANCH})
    content = base64.b64decode(result["content"]).decode("utf-8")
    return content, result["sha"]

def github_create_branch(branch, base_sha):
    return github_request("POST", f"/repos/{GITHUB_REPO}/git/refs", json={"ref": f"refs/heads/{branch}", "sha": base_sha})

def github_commit_file(branch, content, file_sha, task):
    encoded = base64.b64encode(content.encode("utf-8")).decode("ascii")
    return github_request("PUT", f"/repos/{GITHUB_REPO}/contents/{CODE_FILE}", json={"message": f"Nova AI: {task[:60]}", "content": encoded, "sha": file_sha, "branch": branch})

def github_admin_only(message):
    if ADMIN_ID == 0 or message.from_user.id != ADMIN_ID:
        bot.reply_to(message, "⛔ GitHub-управление доступно только администратору.")
        return False
    return True

@bot.message_handler(commands=["code"])
def cmd_code(m):
    if not github_admin_only(m): return
    if not GITHUB_TOKEN:
        bot.reply_to(m, "❌ Добавь GITHUB_TOKEN в Environment на Render.")
        return
    task = m.text.partition(" ")[2].strip()
    if not task:
        bot.reply_to(m, "Использование:\n/code Исправь обработку ошибок в Nova.ai.py")
        return
    if len(task) > 1000:
        bot.reply_to(m, "Слишком длинное задание (максимум 1000 символов).")
        return

    tmp = bot.reply_to(m, "🔎 Получаю Nova.ai.py из GitHub и готовлю изменение...")
    try:
        original, file_sha = github_get_file()
        if len(original) > 45000: raise RuntimeError("Файл слишком большой для безопасного изменения.")

        prompt = (
            "Измени исходный Python-файл по заданию пользователя.\n"
            "Верни ТОЛЬКО корректный JSON без Markdown:\n"
            '{"content":"полный новый исходный код","summary":"краткое описание"}\n'
            "content должен содержать ПОЛНЫЙ файл, не фрагмент.\n"
            "Сохрани существующие функции, если изменение их не требует.\n"
            "Не добавляй секреты и токены. Не изменяй механизм запуска произвольных команд.\n"
            f"ЗАДАНИЕ:\n{task}\n\n"
            f"ТЕКУЩИЙ ФАЙЛ:\n{original}"
        )

        answer = ask_groq(m.chat.id, prompt, system_prompt="Ты — осторожный Python-разработчик. Верни только JSON с полями content и summary. content — полный исходный файл.")
        
        # Очищаем возможные markdown-обертки от модели
        answer = re.sub(r"^\s*```[a-zA-Z0-9_-]*\s*\n?", "", answer)
        answer = re.sub(r"\n?```\s*$", "", answer).strip()
        
        result = json.loads(answer)
        new_content = result.get("content")
        summary = result.get("summary", "Изменение кода")

        if not isinstance(new_content, str) or not new_content.strip(): raise ValueError("Модель не вернула новый код.")
        if len(new_content) > 90000: raise ValueError("Новая версия файла слишком большая.")

        compile(new_content, CODE_FILE, "exec")

        if new_content == original:
            bot.edit_message_text("Изменений нет: новая версия совпадает с текущей.", m.chat.id, tmp.message_id)
            return

        diff = "".join(difflib.unified_diff(original.splitlines(keepends=True), new_content.splitlines(keepends=True), fromfile=f"a/{CODE_FILE}", tofile=f"b/{CODE_FILE}"))

        _, latest_sha = github_get_file()
        if latest_sha != file_sha: raise RuntimeError("Файл изменился в GitHub во время подготовки. Запусти /code ещё раз.")

        change_id = str(time.time_ns())[-10:]
        pending_code_changes[change_id] = {
            "chat_id": m.chat.id, "user_id": m.from_user.id, "content": new_content,
            "file_sha": file_sha, "summary": summary[:500], "task": task[:100],
            "diff": diff, "created": time.time()
        }

        markup = types.InlineKeyboardMarkup()
        markup.add(types.InlineKeyboardButton("✅ Создать ветку и сохранить", callback_data=f"code_yes_{change_id}"), types.InlineKeyboardButton("❌ Отменить", callback_data=f"code_no_{change_id}"))

        bot.edit_message_text(f"🛠 Подготовлено изменение {CODE_FILE}\n\nЗадание: {task[:500]}\n\nОписание: {summary[:500]}\n\nИзменение прошло проверку синтаксиса. Оно ещё не сохранено в GitHub.", m.chat.id, tmp.message_id, reply_markup=markup)
        diff_file = io.BytesIO(diff.encode("utf-8"))
        diff_file.name = "Nova.ai.diff"
        bot.send_document(m.chat.id, diff_file, caption="📄 Проверь diff перед подтверждением.")
    except Exception as e:
        try: bot.edit_message_text(f"❌ Не удалось подготовить изменение:\n{str(e)[:1000]}", m.chat.id, tmp.message_id)
        except: bot.send_message(m.chat.id, f"❌ Ошибка GitHub Code: {str(e)[:1000]}")

@bot.callback_query_handler(func=lambda call: call.data.startswith("code_yes_") or call.data.startswith("code_no_"))
def handle_code_confirmation(call):
    if ADMIN_ID == 0 or not call.from_user or call.from_user.id != ADMIN_ID:
        bot.answer_callback_query(call.id, "Нет доступа.", show_alert=True)
        return

    action, _, change_id = call.data.split("_", 2)
    change = pending_code_changes.get(change_id)

    if not change:
        bot.answer_callback_query(call.id, "Изменение не найдено или уже обработано.", show_alert=True)
        return

    if change["user_id"] != call.from_user.id or change["chat_id"] != call.message.chat.id:
        bot.answer_callback_query(call.id, "Нет доступа.", show_alert=True)
        return

    if time.time() - change["created"] > 900:
        pending_code_changes.pop(change_id, None)
        bot.answer_callback_query(call.id, "Время подтверждения истекло. Запусти /code заново.", show_alert=True)
        return

    if action == "code":
        pending_code_changes.pop(change_id, None)
        bot.answer_callback_query(call.id, "Изменение отменено.")
        bot.edit_message_text("❌ Изменение отменено. GitHub не изменён.", call.message.chat.id, call.message.message_id)
        return

    bot.answer_callback_query(call.id, "Сохраняю изменение...")
    try:
        _, latest_sha = github_get_file()
        if latest_sha != change["file_sha"]:
            raise RuntimeError("Базовый файл изменился. Изменение отменено; подготовь его заново.")

        base_ref = github_request("GET", f"/repos/{GITHUB_REPO}/git/ref/heads/{GITHUB_BASE_BRANCH}")
        branch = "nova-ai/" + datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S") + "-" + change_id

        github_create_branch(branch, base_ref["object"]["sha"])
        result = github_commit_file(branch, change["content"], change["file_sha"], change["task"])
        pending_code_changes.pop(change_id, None)

        bot.edit_message_text(f"✅ Изменение сохранено в отдельную ветку GitHub.\n\nВетка: {branch}\nКоммит: {result.get('commit', {}).get('sha', '')}\n\nОсновная ветка не изменена. Проверь изменения и создай Pull Request.", call.message.chat.id, call.message.message_id)
    except Exception as e:
        bot.send_message(call.message.chat.id, f"❌ Не удалось сохранить изменение:\n{str(e)[:1000]}")

# =========================================================
# NOVA FILE STUDIO — ГЕНЕРАТОР ФАЙЛОВ
# =========================================================

@bot.message_handler(commands=["file"])
def file_menu(m):
    if not is_allowed(m): return
    markup = types.InlineKeyboardMarkup(row_width=1)
    markup.add(types.InlineKeyboardButton("📊 Excel — XLSX", callback_data="filetype_xlsx"), types.InlineKeyboardButton("📽 Презентация — PPTX", callback_data="filetype_pptx"), types.InlineKeyboardButton("📝 Текстовый документ — TXT", callback_data="filetype_txt"))
    bot.reply_to(m, "📁 Nova File Studio\n\nВыбери формат файла:", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("filetype_"))
def choose_file_type(call):
    if not is_allowed(call.message): return
    file_type = call.data.replace("filetype_", "", 1)
    if file_type not in ("xlsx", "pptx", "txt"): return
    bot.answer_callback_query(call.id)
    prompt_msg = bot.send_message(call.message.chat.id, "✍️ Опиши, какой файл создать.\n\nПримеры:\n• Таблица расходов за месяц\n• Презентация об искусственном интеллекте\n• Документ о космосе")
    bot.register_next_step_handler(prompt_msg, process_file_request, file_type)

def clean_model_output(text):
    text = text.strip()
    text = re.sub(r"^\s*```[a-zA-Z0-9_-]*\s*\n?", "", text)
    text = re.sub(r"\n?```\s*$", "", text)
    return text.strip()

def process_file_request(m, file_type):
    if not is_allowed(m): return
    if not m.text or not m.text.strip(): return
    if len(m.text) > 2000:
        bot.reply_to(m, "Описание слишком длинное. Сократи его до 2000 символов.")
        return

    tmp = bot.reply_to(m, "⚙️ Nova создаёт файл...")
    try:
        if file_type == "txt":
            content = ask_groq(m.chat.id, "Создай текстовый документ по запросу. Верни только содержимое.\n\n" + m.text, system_prompt="Ты — редактор документов.")
            output = io.BytesIO(content.encode("utf-8"))
            output.name = "nova_document.txt"

        elif file_type == "xlsx":
            content = ask_groq(m.chat.id, "Создай таблицу. Верни ТОЛЬКО CSV с разделителем (;).\n\n" + m.text, system_prompt="Ты создаёшь таблицы для Excel.")
            content = clean_model_output(content)
            rows = list(csv.reader(io.StringIO(content), delimiter=";"))
            rows = [[cell.strip() for cell in row] for row in rows if any(cell.strip() for cell in row)]
            if not rows: raise ValueError("ИИ не создал таблицу.")
            max_columns = max(len(row) for row in rows)
            rows = [row + [""] * (max_columns - len(row)) for row in rows]

            wb = Workbook()
            ws = wb.active
            ws.title = "Nova SI"
            for row in rows: ws.append(row)

            header_fill = PatternFill("solid", fgColor="2457A7")
            alternate_fill = PatternFill("solid", fgColor="EAF2FF")
            thin_border = Border(bottom=Side(style="thin", color="D6E0F0"))

            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 30

            for row_index in range(2, ws.max_row + 1):
                for cell in ws[row_index]:
                    cell.font = Font(name="Arial", size=10, color="263449")
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.border = thin_border
                    if row_index % 2 == 0: cell.fill = alternate_fill

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False

            for column in ws.columns:
                letter = column[0].column_letter
                max_length = max(len(str(cell.value or "")) for cell in column)
                ws.column_dimensions[letter].width = min(max(max_length + 3, 14), 35)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            output.name = "nova_table.xlsx"

        elif file_type == "pptx":
            content = ask_groq(m.chat.id, "Создай презентацию на 6 слайдов. Формат: 'СЛАЙД: Заголовок' и тезисы с '-'. Без Markdown.\n\n" + m.text, system_prompt="Ты — автор презентаций.")
            content = clean_model_output(content)

            sections = re.split(r"(?im)^\s*СЛАЙД\s*:\s*", content)
            slides_data = []
            for section in sections:
                lines = [line.strip() for line in section.splitlines() if line.strip()]
                if not lines: continue
                title = re.sub(r"^\d+[\.\):\-]\s*", "", lines[0].lstrip("#").strip())
                bullets = [re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip() for line in lines[1:] if line.strip()]
                slides_data.append((title[:120], bullets[:5]))

            if not slides_data: raise ValueError("ИИ не вернул слайды.")

            slides_data = slides_data[:8]
            prs = Presentation()
            prs.slide_width = Inches(12)
            prs.slide_height = Inches(6.75)
            NAVY = RGBColor(18, 29, 54); BLUE = RGBColor(48, 110, 230); WHITE = RGBColor(255, 255, 255); LIGHT = RGBColor(229, 237, 250); MUTED = RGBColor(170, 190, 220)

            for index, (title, bullets) in enumerate(slides_data, start=1):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                background = slide.background.fill
                background.solid()
                background.fore_color.rgb = NAVY

                accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(12), Inches(0.12))
                accent.fill.solid(); accent.fill.fore_color.rgb = BLUE; accent.line.fill.background()

                number_box = slide.shapes.add_textbox(Inches(10.75), Inches(0.35), Inches(0.7), Inches(0.35))
                number_p = number_box.text_frame.paragraphs[0]
                number_p.text = f"{index:02d} / {len(slides_data):02d}"
                number_p.alignment = PP_ALIGN.RIGHT; number_p.font.size = Pt(10); number_p.font.color.rgb = MUTED

                title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(10.1), Inches(1.15))
                title_tf = title_box.text_frame; title_tf.clear(); title_tf.word_wrap = True; title_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                title_p = title_tf.paragraphs[0]
                title_p.text = title or f"Слайд {index}"
                title_p.font.name = "Arial"; title_p.font.size = Pt(30 if len(title) < 45 else 24); title_p.font.bold = True; title_p.font.color.rgb = WHITE

                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.95), Inches(10.6), Inches(3.95))
                card.fill.solid(); card.fill.fore_color.rgb = RGBColor(28, 43, 73); card.line.color.rgb = RGBColor(48, 67, 103)

                body_box = slide.shapes.add_textbox(Inches(1.05), Inches(2.18), Inches(9.9), Inches(3.45))
                body_tf = body_box.text_frame; body_tf.clear(); body_tf.word_wrap = True; body_tf.vertical_anchor = MSO_ANCHOR.MIDDLE

                bullets = bullets if bullets else ["Основная информация по теме."]
                for bullet_index, bullet in enumerate(bullets[:5]):
                    paragraph = body_tf.paragraphs[0] if bullet_index == 0 else body_tf.add_paragraph()
                    paragraph.text = f"•  {bullet}"
                    paragraph.font.name = "Arial"; paragraph.font.size = Pt(19 if len(bullets) <= 4 else 16); paragraph.font.color.rgb = LIGHT; paragraph.space_after = Pt(15)

                footer_box = slide.shapes.add_textbox(Inches(0.75), Inches(6.15), Inches(10.5), Inches(0.3))
                footer_p = footer_box.text_frame.paragraphs[0]
                footer_p.text = "NOVA SI  •  PRESENTATION"
                footer_p.font.name = "Arial"; footer_p.font.size = Pt(9); footer_p.font.color.rgb = MUTED

            output = io.BytesIO()
            prs.save(output)
            output.seek(0)
            output.name = "nova_presentation.pptx"

        bot.send_document(m.chat.id, output, caption="✅ Файл создан — Nova SI")
        bot.delete_message(m.chat.id, tmp.message_id)
    except Exception as e:
        try: bot.edit_message_text(f"❌ Не получилось создать файл: {str(e)[:800]}", m.chat.id, tmp.message_id)
        except: bot.send_message(m.chat.id, f"❌ Ошибка файла: {str(e)[:800]}")

# =========================================================
# NOVA CORE AUTO
# =========================================================

AGENT_PROMPTS = {
    "planner": "Ты — Planner (Планировщик) Nova Core. Разбей задачу на пошаговые инструкции. Нумерованный список. Без воды.",
    "builder": "Ты — Builder (Разработчик). Напиши решение по плану. Чистый Python-код.",
    "checker": "Ты — Checker (Проверяющий). Проверь решение на баги, уязвимости и ошибки. Будь строгим."
}

@bot.message_handler(commands=["auto"])
def cmd_auto(m):
    if not is_allowed(m): return
    msg = bot.reply_to(m, "🚀 Nova Core Auto активирован.\n\nНапиши задачу. Агенты выполнят её по цепочке (План → Решение → Проверка).")
    bot.register_next_step_handler(msg, process_auto_task)

def process_auto_task(m):
    if not is_allowed(m): return
    task_text = m.text
    if not task_text or not task_text.strip(): return

    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("INSERT INTO core_tasks (chat_id, task_text) VALUES (?, ?)", (m.chat.id, task_text))
    task_id = c.lastrowid
    conn.commit()
    conn.close()

    bot.reply_to(m, f"Задача: {task_text[:100]}\n\nЗапускаю цепочку агентов.")

    def safe_groq(prompt, sys_prompt):
        safe_prompt = prompt[:3000]
        for attempt in range(3):
            try: return ask_groq(m.chat.id, safe_prompt, system_prompt=sys_prompt)
            except Exception as e:
                if "429" in str(e):
                    bot.send_message(m.chat.id, "⏳ Лимит Groq. Жду 30 секунд.")
                    time.sleep(30)
                elif "413" in str(e): safe_prompt = safe_prompt[:1500]
                else: raise
        raise Exception("Не удалось получить ответ от Groq.")

    try:
        msg1 = bot.send_message(m.chat.id, "⏳ Шаг 1: Planner составляет план...")
        plan = safe_groq(task_text, AGENT_PROMPTS["planner"])
        bot.edit_message_text("✅ План готов. Следующий шаг...", m.chat.id, msg1.message_id)
        time.sleep(3)

        msg2 = bot.send_message(m.chat.id, "⏳ Шаг 2: Builder готовит решение...")
        builder_input = f"Задача:\n{task_text}\n\nПлан:\n{plan[:2000]}\n\nПодготовь решение по этому плану."
        code_result = safe_groq(builder_input, AGENT_PROMPTS["builder"])
        bot.edit_message_text("✅ Решение готово. Проверяю...", m.chat.id, msg2.message_id)
        time.sleep(3)

        msg3 = bot.send_message(m.chat.id, "⏳ Шаг 3: Checker проверяет решение...")
        checker_input = f"Задача:\n{task_text}\n\nРешение:\n{code_result[:2000]}\n\nПроверь решение."
        review = safe_groq(checker_input, AGENT_PROMPTS["checker"])

        final_result = f"ПЛАН:\n{plan}\n\nРЕШЕНИЕ:\n{code_result}\n\nПРОВЕРКА:\n{review}"
        conn = sqlite3.connect("bot.db")
        c = conn.cursor()
        c.execute("UPDATE core_tasks SET agent_type = 'auto', result = ?, status = 'done' WHERE id = ?", (final_result, task_id))
        conn.commit()
        conn.close()

        bot.edit_message_text("✅ Проверка завершена!", m.chat.id, msg3.message_id)
        final_text = f"🔥 РЕЗУЛЬТАТ NOVA CORE AUTO\n\n📋 ПЛАН:\n{plan}\n\n💻 РЕШЕНИЕ:\n{code_result}\n\n🔍 ПРОВЕРКА:\n{review}"
        for i in range(0, len(final_text), 4000): bot.send_message(m.chat.id, final_text[i:i + 4000])
    except Exception as e:
        bot.send_message(m.chat.id, f"❌ Ошибка в цепочке: {str(e)[:800]}")

# =========================================================
# NOVA CORE — ОТДЕЛЬНЫЕ АГЕНТЫ
# =========================================================

@bot.message_handler(commands=["core"])
def cmd_core(m):
    if not is_allowed(m): return
    msg = bot.reply_to(m, "🧠 Nova Core активирован.\n\nНапиши задачу:")
    bot.register_next_step_handler(msg, process_core_task)

def process_core_task(m):
    if not is_allowed(m): return
    task_text = m.text
    if not task_text or not task_text.strip(): return
    
    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("INSERT INTO core_tasks (chat_id, task_text) VALUES (?, ?)", (m.chat.id, task_text))
    task_id = c.lastrowid
    conn.commit()
    conn.close()

    markup = types.InlineKeyboardMarkup(row_width=1)
    markup.add(types.InlineKeyboardButton("📋 Planner", callback_data=f"core_planner_{task_id}"), types.InlineKeyboardButton("💻 Builder", callback_data=f"core_builder_{task_id}"), types.InlineKeyboardButton("🔍 Checker", callback_data=f"core_checker_{task_id}"))
    bot.reply_to(m, f"Задача: {task_text[:100]}\n\nВыбери агента:", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("core_"))
def handle_core_callback(call):
    if not is_allowed(call.message): return
    parts = call.data.split("_")
    if len(parts) != 3: return
    agent_type = parts[1]
    try: task_id = int(parts[2])
    except ValueError: return

    if agent_type not in AGENT_PROMPTS: return
    bot.answer_callback_query(call.id, f"Агент {agent_type} работает...")

    conn = sqlite3.connect("bot.db")
    c = conn.cursor()
    c.execute("SELECT task_text FROM core_tasks WHERE id = ? AND chat_id = ?", (task_id, call.message.chat.id))
    row = c.fetchone()
    conn.close()
    if not row: return

    task_text = row[0]
    try: bot.edit_message_text(f"⏳ Агент {agent_type.capitalize()} думает...", call.message.chat.id, call.message.message_id)
    except: pass

    try:
        answer = ask_groq(call.message.chat.id, task_text[:3000], system_prompt=AGENT_PROMPTS[agent_type])
        conn = sqlite3.connect("bot.db")
        c = conn.cursor()
        c.execute("UPDATE core_tasks SET agent_type = ?, result = ?, status = 'done' WHERE id = ? AND chat_id = ?", (agent_type, answer, task_id, call.message.chat.id))
        conn.commit()
        conn.close()

        full_text = f"🤖 Ответ от {agent_type.capitalize()}:\n\n{answer}"
        chunks = [full_text[i:i + 4000] for i in range(0, len(full_text), 4000)]
        for index, chunk in enumerate(chunks):
            if index == 0:
                try: bot.edit_message_text(chunk, call.message.chat.id, call.message.message_id)
                except: bot.send_message(call.message.chat.id, chunk)
            else: bot.send_message(call.message.chat.id, chunk)
    except Exception as e:
        bot.send_message(call.message.chat.id, f"❌ Ошибка агента: {str(e)[:800]}")

# =========================================================
# КОМАНДЫ
# =========================================================

@bot.message_handler(commands=["start"])
def start(m):
    if not is_allowed(m): return
    get_user(m.chat.id)
    bot.send_message(m.chat.id, f"👋 Привет, {m.from_user.first_name}!\n\nЯ — Nova SI. Выбери раздел ниже:", reply_markup=get_main_inline_menu())
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    markup.add(types.KeyboardButton("Очистить чат"), types.KeyboardButton("Сменить модель"))
    bot.send_message(m.chat.id, "Кнопки быстрого доступа:", reply_markup=markup)

@bot.message_handler(commands=["help"])
def cmd_help(m):
    if not is_allowed(m): return
    help_text = "🤖 Возможности Nova SI:\n\n💬 ИИ-чат — просто напиши.\n📎 Файлы — PDF, TXT, картинки.\n🖼 Картинки — /image запрос.\n🎙 Голос — отправь голосовое.\n🌐 Поиск — кнопка Поиск.\n🚀 Nova Core Auto — /auto.\n🧠 Отдельные агенты — /core.\n📁 Генератор файлов — /file.\n🛠 Код в GitHub — /code."
    bot.send_message(m.chat.id, help_text)

@bot.message_handler(commands=["clear"])
def cmd_clear(m):
    if not is_allowed(m): return
    clear_history(m.chat.id)
    bot.reply_to(m, "🧹 История очищена.", reply_markup=get_main_keyboard(m.chat.id))

@bot.message_handler(commands=["profile"])
def cmd_profile(m):
    if not is_allowed(m): return
    role, model, requests_count = get_user(m.chat.id)
    bot.reply_to(m, f"👤 Профиль:\n\nРоль: {role}\nМодель: {model}\nЗапросов: {requests_count}", reply_markup=get_main_keyboard(m.chat.id))

# =========================================================
# ГЕНЕРАЦИЯ ИЗОБРАЖЕНИЙ
# =========================================================

@bot.message_handler(commands=["image"])
def gen_image(m):
    if not is_allowed(m): return
    prompt = m.text.replace("/image", "", 1).strip()
    if not prompt:
        bot.reply_to(m, "Напиши описание изображения после /image.")
        return
    tmp = bot.reply_to(m, "🎨 Создаю изображение...")
    try:
        url = f"https://image.pollinations.ai/prompt/{requests.utils.quote(prompt)}?width=1024&height=1024&nologo=true"
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        if not response.content: raise ValueError("Сервис вернул пустое изображение.")
        bot.delete_message(m.chat.id, tmp.message_id)
        photo = io.BytesIO(response.content)
        photo.name = "nova_image.jpg"
        bot.send_photo(m.chat.id, photo, caption=f"🎨 Запрос: {prompt[:800]}")
    except Exception as e:
        bot.edit_message_text(f"Ошибка создания изображения: {str(e)[:700]}", m.chat.id, tmp.message_id)

# =========================================================
# ГОЛОСОВЫЕ СООБЩЕНИЯ
# =========================================================

@bot.message_handler(content_types=["voice"])
def handle_voice(m):
    if not is_allowed(m): return
    tmp = bot.reply_to(m, "🎙 Слушаю...")
    audio_file = f"/tmp/ans_{m.chat.id}.mp3"
    try:
        file_info = bot.get_file(m.voice.file_id)
        downloaded = bot.download_file(file_info.file_path)
        url = "https://api.groq.com/openai/v1/audio/transcriptions"
        headers = {"Authorization": f"Bearer {GROQ_KEY}"}
        files = {"file": ("voice.ogg", downloaded)}
        data = {"model": "whisper-large-v3-turbo"}
        response = requests.post(url, headers=headers, files=files, data=data, timeout=60)
        response.raise_for_status()
        text = response.json().get("text", "").strip()
        if not text: raise ValueError("Не удалось распознать речь.")
        bot.edit_message_text(f"Ты сказал: {text[:1000]}\n\nДумаю...", m.chat.id, tmp.message_id)
        answer = ask_groq(m.chat.id, text)
        async def tts():
            communicate = edge_tts.Communicate(answer[:4000], "ru-RU-DmitryNeural")
            await communicate.save(audio_file)
        asyncio.run(tts())
        with open(audio_file, "rb") as audio: bot.send_voice(m.chat.id, audio, caption=answer[:900])
        os.remove(audio_file)
    except Exception as e:
        bot.edit_message_text(f"Ошибка обработки голоса: {str(e)[:700]}", m.chat.id, tmp.message_id)
    finally:
        if os.path.exists(audio_file):
            try: os.remove(audio_file)
            except OSError: pass

# =========================================================
# ПОИСК В ИНТЕРНЕТЕ
# =========================================================

@bot.message_handler(func=lambda message: message.text == "Поиск")
def search_web(m):
    if not is_allowed(m): return
    msg = bot.reply_to(m, "Введи поисковый запрос:")
    bot.register_next_step_handler(msg, process_search)

def process_search(m):
    if not is_allowed(m): return
    if not m.text or not m.text.strip(): return
    tmp = bot.reply_to(m, "🌐 Ищу в интернете...")
    instances = ["https://searx.tiekoetter.com", "https://searx.be", "https://searxng.site", "https://priv.au"]
    headers = {"User-Agent": "Mozilla/5.0"}
    results = None
    for base_url in instances:
        try:
            response = requests.get(f"{base_url}/search", params={"q": m.text, "format": "json"}, headers=headers, timeout=10)
            if response.status_code == 200:
                data = response.json()
                if data.get("results"): results = data["results"][:3]; break
        except: continue
    if not results:
        bot.edit_message_text("Поисковые сервисы сейчас недоступны. Попробуй позже.", m.chat.id, tmp.message_id)
        return
    context = "".join(f"Источник: {result.get('title', '')}\nСсылка: {result.get('url', '')}\n{result.get('content', '')}\n\n" for result in results)
    try:
        answer = ask_groq(m.chat.id, f"Ответь на запрос: {m.text}\n\nИспользуй информацию:\n{context[:3000]}")
        bot.edit_message_text(answer[:4000], m.chat.id, tmp.message_id)
    except Exception as e:
        bot.edit_message_text(f"Ошибка ИИ: {str(e)[:700]}", m.chat.id, tmp.message_id)

# =========================================================
# ИНЛАЙН-МЕНЮ
# =========================================================

@bot.callback_query_handler(func=lambda call: call.data.startswith("menu_"))
def handle_menu(call):
    if not is_allowed(call.message): return
    bot.answer_callback_query(call.id)
    data = call.data
    if data == "menu_chat": bot.edit_message_text("💬 ИИ-Чат активирован. Отправь сообщение.", call.message.chat.id, call.message.message_id, reply_markup=get_main_inline_menu())
    elif data == "menu_modes":
        markup = types.InlineKeyboardMarkup(row_width=2)
        for role in ["Обычный", "Дебаггер", "Эксперт", "Юморист", "Программист", "Дружелюбный"]:
            markup.add(types.InlineKeyboardButton(role, callback_data=f"role_{role}"))
        markup.add(types.InlineKeyboardButton("🔙 Назад", callback_data="menu_back"))
        bot.edit_message_text("🧠 Выбери режим:", call.message.chat.id, call.message.message_id, reply_markup=markup)
    elif data == "menu_history":
        history = get_history(call.message.chat.id)
        bot.edit_message_text("📜 Последние сообщения в истории." if history else "📜 История пуста.", call.message.chat.id, call.message.message_id, reply_markup=get_main_inline_menu())
    elif data == "menu_profile":
        role, model, reqs = get_user(call.message.chat.id)
        bot.edit_message_text(f"👤 Профиль\n\nРоль: {role}\nМодель: {model}\nЗапросов: {reqs}", call.message.chat.id, call.message.message_id, reply_markup=get_main_inline_menu())
    elif data == "menu_settings": bot.edit_message_text("⚙️ Настройки\n\nДля выбора модели используй кнопку «Сменить модель».", call.message.chat.id, call.message.message_id, reply_markup=get_main_inline_menu())
    elif data == "menu_help": bot.edit_message_text("🤖 Nova SI\n\n/start — меню\n/help — помощь\n/clear — очистить\n/profile — профиль\n/file — файлы\n/image — картинки\n/core — агенты\n/auto — цепочка\n/code — GitHub", call.message.chat.id, call.message.message_id, reply_markup=get_main_inline_menu())
    elif data == "menu_back": bot.edit_message_text("Выбери раздел:", call.message.chat.id, call.message.message_id, reply_markup=get_main_inline_menu())

# =========================================================
# ВЫБОР РОЛИ
# =========================================================

@bot.message_handler(func=lambda message: message.text.startswith("Роль:"))
def role_menu(m):
    if not is_allowed(m): return
    markup = types.InlineKeyboardMarkup()
    for role in ["Обычный", "Дебаггер", "Эксперт", "Юморист", "Программист", "Дружелюбный"]:
        markup.add(types.InlineKeyboardButton(role, callback_data=f"role_{role}"))
    bot.reply_to(m, "Выбери роль:", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("role_"))
def set_role(call):
    if not is_allowed(call.message): return
    role = call.data.replace("role_", "", 1)
    allowed_roles = {"Обычный", "Дебаггер", "Эксперт", "Юморист", "Программист", "Дружелюбный"}
    if role not in allowed_roles: return
    update_user(call.message.chat.id, role=role)
    bot.answer_callback_query(call.id, f"Роль: {role}")
    try: bot.edit_message_text(f"Теперь моя роль: {role}", call.message.chat.id, call.message.message_id)
    except: pass
    bot.send_message(call.message.chat.id, "Меню обновлено.", reply_markup=get_main_keyboard(call.message.chat.id))

# =========================================================
# СМЕНА МОДЕЛИ
# =========================================================

@bot.message_handler(func=lambda message: message.text == "Сменить модель")
def model_menu(m):
    if not is_allowed(m): return
    tmp = bot.reply_to(m, "Смотрю доступные модели...")
    try:
        response = requests.get("https://api.groq.com/openai/v1/models", headers={"Authorization": f"Bearer {GROQ_KEY}"}, timeout=20)
        response.raise_for_status()
        valid_models = [item["id"] for item in response.json().get("data", []) if not any(word in item["id"].lower() for word in ["whisper", "guard", "tts"])]
        if not valid_models: return
        markup = types.InlineKeyboardMarkup()
        for model_id in valid_models[:10]: markup.add(types.InlineKeyboardButton(model_id, callback_data=f"model_{model_id}"))
        bot.edit_message_text("Выбери модель:", m.chat.id, tmp.message_id, reply_markup=markup)
    except Exception as e:
        bot.edit_message_text(f"Ошибка получения моделей: {str(e)[:700]}", m.chat.id, tmp.message_id)

@bot.callback_query_handler(func=lambda call: call.data.startswith("model_"))
def set_model(call):
    if not is_allowed(call.message): return
    model_id = call.data.replace("model_", "", 1)
    update_user(call.message.chat.id, model=model_id)
    bot.answer_callback_query(call.id, "Модель обновлена.")
    bot.edit_message_text(f"Теперь я работаю на модели:\n{model_id}", call.message.chat.id, call.message.message_id)

@bot.message_handler(func=lambda message: message.text == "Профиль")
def profile_btn(m): cmd_profile(m)

@bot.message_handler(func=lambda message: message.text in ["Очистить историю", "Очистить чат"])
def clear_btn(m): cmd_clear(m)

# =========================================================
# ОБРАБОТКА ФОТОГРАФИЙ
# =========================================================

@bot.message_handler(content_types=["photo"])
def handle_photo(m):
    if not is_allowed(m): return
    if not OPENROUTER_KEY:
        bot.reply_to(m, "Не задан ключ OPENROUTER_KEY.")
        return
    tmp = bot.reply_to(m, "🖼 Анализирую изображение...")
    try:
        file_info = bot.get_file(m.photo[-1].file_id)
        image_bytes = bot.download_file(file_info.file_path)
        base64_image = base64.b64encode(image_bytes).decode("utf-8")
        url = "https://openrouter.ai/api/v1/chat/completions"
        headers = {"Authorization": f"Bearer {OPENROUTER_KEY}", "Content-Type": "application/json"}
        recognized_text = None
        for model_id in ["google/gemma-4-26b-a4b-it:free", "nvidia/nemotron-nano-12b-v2-vl:free"]:
            payload = {"model": model_id, "messages": [{"role": "user", "content": [{"type": "text", "text": "Распознай текст на изображении."}, {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}]}]}
            try:
                response = requests.post(url, headers=headers, json=payload, timeout=40)
                if response.status_code == 200:
                    recognized_text = response.json()["choices"][0]["message"]["content"]
                    if recognized_text: break
            except: continue
        if not recognized_text:
            bot.edit_message_text("Модели распознавания недоступны.", m.chat.id, tmp.message_id)
            return
        prompt = f"Найди ошибки в решении:\n{recognized_text[:2500]}" if find_error_mode.get(m.chat.id, False) else f"Реши пошагово:\n{recognized_text[:2500]}"
        answer = ask_groq(m.chat.id, prompt)
        bot.delete_message(m.chat.id, tmp.message_id)
        send_smart_response(m.chat.id, answer)
    except Exception as e:
        try: bot.edit_message_text(f"Ошибка анализа: {str(e)[:700]}", m.chat.id, tmp.message_id)
        except: pass

# =========================================================
# РЕЖИМ «НАЙДИ ОШИБКУ»
# =========================================================

@bot.message_handler(func=lambda message: message.text.startswith("Найди ошибку:"))
def toggle_error_mode(m):
    if not is_allowed(m): return
    current = find_error_mode.get(m.chat.id, False)
    find_error_mode[m.chat.id] = not current
    status = "ВКЛ" if not current else "ВЫКЛ"
    bot.reply_to(m, f"Режим поиска ошибок: {status}", reply_markup=get_main_keyboard(m.chat.id))

# =========================================================
# ПРИЁМ ДОКУМЕНТОВ И КОНВЕРТАЦИЯ
# =========================================================

@bot.message_handler(content_types=["document"])
def handle_doc(m):
    if not is_allowed(m): return
    file_name = m.document.file_name or "uploaded_file"
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    try:
        file_info = bot.get_file(m.document.file_id)
        downloaded = bot.download_file(file_info.file_path)
    except Exception as e:
        bot.reply_to(m, f"Не удалось скачать файл: {str(e)[:500]}")
        return
    user_files[m.chat.id] = {"name": file_name, "data": downloaded, "ext": ext}
    if ext in ["pdf", "txt"]:
        tmp = bot.reply_to(m, "📄 Читаю документ...")
        try:
            if ext == "pdf":
                pages = PdfReader(io.BytesIO(downloaded)).pages
                text = "\n".join(page.extract_text() or "" for page in pages)
            else:
                text = downloaded.decode("utf-8", errors="ignore")
            if not text.strip(): return
            answer = ask_groq(m.chat.id, f"Проанализируй документ:\n\n{text[:3000]}")
            bot.edit_message_text(answer[:4000], m.chat.id, tmp.message_id)
            user_files.pop(m.chat.id, None)
        except Exception as e:
            bot.edit_message_text(f"Ошибка чтения: {str(e)[:700]}", m.chat.id, tmp.message_id)
        return
    markup = types.InlineKeyboardMarkup()
    if ext == "pdf": markup.add(types.InlineKeyboardButton("В DOCX", callback_data="conv_pdf_docx"))
    elif ext in ["jpg", "jpeg", "png", "webp"]: markup.add(types.InlineKeyboardButton("В PDF", callback_data="conv_img_pdf"))
    elif ext == "csv": markup.add(types.InlineKeyboardButton("В Excel", callback_data="conv_csv_xlsx"))
    elif ext in ["mp3", "ogg", "wav", "m4a"]: markup.add(types.InlineKeyboardButton("В текст", callback_data="conv_audio_txt"))
    else:
        user_files.pop(m.chat.id, None)
        bot.reply_to(m, f"Формат .{ext or '?'} не поддерживается.")
        return
    bot.reply_to(m, f"Файл «{file_name}» получен. Что делаем?", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("conv_"))
def process_conversion(call):
    if not is_allowed(call.message): return
    chat_id = call.message.chat.id
    if chat_id not in user_files: return
    bot.answer_callback_query(call.id)
    file_data = user_files[chat_id]
    tmp = bot.send_message(chat_id, "⚙️ Конвертирую...")
    temp_paths = []
    try:
        if call.data == "conv_pdf_docx":
            source_path = "/tmp/nova_source.pdf"
            output_path = "/tmp/nova_converted.docx"
            temp_paths.extend([source_path, output_path])
            with open(source_path, "wb") as file: file.write(file_data["data"])
            converter = pdf2docx.Converter(source_path)
            try: converter.convert(output_path)
            finally: converter.close()
            with open(output_path, "rb") as file: document = io.BytesIO(file.read())
            document.name = "nova_converted.docx"
            bot.send_document(chat_id, document, caption="✅ Готово: DOCX")
        elif call.data == "conv_img_pdf":
            img = Image.open(io.BytesIO(file_data["data"]))
            pdf_bytes = io.BytesIO()
            img.convert("RGB").save(pdf_bytes, format="PDF")
            pdf_bytes.seek(0)
            pdf_bytes.name = "nova_image.pdf"
            bot.send_document(chat_id, pdf_bytes, caption="✅ Готово: PDF")
        elif call.data == "conv_csv_xlsx":
            df = pd.read_csv(io.BytesIO(file_data["data"]))
            xlsx_bytes = io.BytesIO()
            df.to_excel(xlsx_bytes, index=False)
            xlsx_bytes.seek(0)
            xlsx_bytes.name = "nova_data.xlsx"
            bot.send_document(chat_id, xlsx_bytes, caption="✅ Готово: Excel")
        elif call.data == "conv_audio_txt":
            response = requests.post("https://api.groq.com/openai/v1/audio/transcriptions", headers={"Authorization": f"Bearer {GROQ_KEY}"}, files={"file": (file_data["name"], file_data["data"])}, data={"model": "whisper-large-v3-turbo"}, timeout=60)
            response.raise_for_status()
            transcript = response.json().get("text", "Пусто.")
            txt_bytes = io.BytesIO(transcript.encode("utf-8"))
            txt_bytes.name = "nova_transcript.txt"
            bot.send_document(chat_id, txt_bytes, caption="✅ Готово: текст")
        bot.delete_message(chat_id, tmp.message_id)
        user_files.pop(chat_id, None)
    except Exception as e:
        try: bot.edit_message_text(f"Ошибка конвертации: {str(e)[:700]}", chat_id, tmp.message_id)
        except: pass
    finally:
        for path in temp_paths:
            if os.path.exists(path):
                try: os.remove(path)
                except: pass

# =========================================================
# ОБРАБОТКА ТЕКСТА
# =========================================================

@bot.message_handler(content_types=["text"])
def on_text(m):
    if not is_allowed(m): return
    if not m.text: return
    if m.text.startswith("Роль:"): return role_menu(m)
    if m.text.startswith("Найди ошибку:"): return toggle_error_mode(m)
    if m.text == "Сменить модель": return model_menu(m)
    if m.text == "Профиль": return profile_btn(m)
    if m.text in ["Очистить историю", "Очистить чат"]: return clear_btn(m)
    if m.text == "Поиск": return search_web(m)
    if len(m.text) > 3000:
        bot.reply_to(m, "❌ Слишком длинный запрос. Сократи его до 3000 символов.")
        return
    tmp = bot.reply_to(m, "💭 Думаю...")
    try:
        answer = ask_groq(m.chat.id, m.text)
        try: bot.delete_message(m.chat.id, tmp.message_id)
        except: pass
        send_smart_response(m.chat.id, answer)
    except Exception as e:
        error = str(e)
        if "413" in error: error_text = "❌ Запрос слишком большой."
        elif "429" in error: error_text = "⏳ Лимит запросов к модели."
        else: error_text = f"Ошибка ИИ: {error[:700]}"
        try: bot.edit_message_text(error_text, m.chat.id, tmp.message_id)
        except: bot.send_message(m.chat.id, error_text)

# =========================================================
# NOVA SI — МОНТАЖ ВИДЕО ПО ОБЫЧНОМУ ОПИСАНИЮ
# =========================================================

montage_requests = {}


def is_montage_request(text):
    """Лёгкое распознавание запроса на монтаж без slash-команды."""
    text = (text or "").lower()
    words = (
        "монтаж", "смонтируй", "смонтировать", "эдит", "edit",
        "tiktok", "тикток", "reels", "рилс", "нарежь видео",
        "сделай клип", "сделай ролик", "переходы", "обработай видео",
        "обработай ролик", "видео для"
    )
    return any(word in text for word in words)


@bot.message_handler(content_types=["text"], func=lambda m: bool(m.text) and is_montage_request(m.text))
def nova_montage_instruction(m):
    if not is_allowed(m):
        return
    montage_requests[m.chat.id] = m.text.strip()[:1500]
    bot.reply_to(
        m,
        "🎬 Задание на монтаж сохранено.\n\n"
        "Теперь отправь видео как обычное видео в Telegram.\n"
        "Я применю доступные эффекты по твоему описанию.\n\n"
        "Поддерживается один ролик до 20 МБ и до 120 секунд."
    )


def _build_montage_filter(instruction):
    text = instruction.lower()
    filters = [
        "scale=720:1280:force_original_aspect_ratio=increase",
        "crop=720:1280",
        "fps=30"
    ]
    if any(x in text for x in ("чёрно-бел", "черно-бел", "black and white", "black&white")):
        filters.append("hue=s=0")
    elif any(x in text for x in ("ярк", "цвет", "cinematic", "кинематограф", "насыщ")):
        filters.append("eq=contrast=1.10:saturation=1.20:brightness=0.02")
    if any(x in text for x in ("замедли", "slow motion", "слоумо", "slow-mo")):
        filters.append("setpts=1.25*PTS")
    if any(x in text for x in ("ускор", "быстрее", "fast motion")):
        filters.append("setpts=0.75*PTS")
    filters.append("format=yuv420p")
    return ",".join(filters)


@bot.message_handler(content_types=["video"])
def nova_montage_video(m):
    if not is_allowed(m):
        return
    instruction = montage_requests.get(m.chat.id)
    if not instruction:
        bot.reply_to(
            m,
            "🎬 Сначала напиши, какой монтаж нужен, например: "
            "«Сделай вертикальный эдит для TikTok, яркие цвета и замедление», "
            "а потом отправь видео ещё раз."
        )
        return
    if m.video.file_size and m.video.file_size > 20 * 1024 * 1024:
        bot.reply_to(m, "❌ Видео больше 20 МБ. Отправь ролик меньшего размера.")
        return
    if not shutil.which("ffmpeg"):
        bot.reply_to(m, "❌ На Railway не найден FFmpeg. Проверь nixpacks.toml и пересобери проект.")
        return

    status = bot.reply_to(m, "🎞️ Nova SI обрабатывает видео. Подожди немного…")
    try:
        file_info = bot.get_file(m.video.file_id)
        video_bytes = bot.download_file(file_info.file_path)
        with tempfile.TemporaryDirectory(prefix="nova_montage_") as workdir:
            input_path = os.path.join(workdir, "input.mp4")
            output_path = os.path.join(workdir, "nova_montage.mp4")
            with open(input_path, "wb") as f:
                f.write(video_bytes)

            command = [
                "ffmpeg", "-y", "-i", input_path,
                "-vf", _build_montage_filter(instruction),
                "-t", "120", "-map", "0:v:0", "-map", "0:a?",
                "-c:v", "libx264", "-preset", "veryfast", "-crf", "23",
                "-c:a", "aac", "-b:a", "128k", "-movflags", "+faststart",
                output_path
            ]
            result = subprocess.run(command, capture_output=True, text=True, timeout=300)
            if result.returncode != 0 or not os.path.isfile(output_path):
                raise RuntimeError((result.stderr or "FFmpeg error")[-1200:])
            if os.path.getsize(output_path) > 49 * 1024 * 1024:
                bot.send_message(m.chat.id, "❌ Результат больше лимита отправки. Попробуй более короткое видео.")
                return
            with open(output_path, "rb") as f:
                bot.send_video(m.chat.id, f, caption="🎬 Готово! Nova SI обработал твой ролик.")
        montage_requests.pop(m.chat.id, None)
        try:
            bot.delete_message(m.chat.id, status.message_id)
        except Exception:
            pass
    except subprocess.TimeoutExpired:
        bot.reply_to(m, "⏳ Обработка заняла слишком много времени. Попробуй короткий ролик.")
    except Exception as e:
        print("Ошибка монтажа:", repr(e))
        bot.reply_to(m, "❌ Не удалось обработать видео. Проверь логи Railway и попробуй другой ролик.")
    finally:
        # Не удаляем задание при технической ошибке: можно повторить отправку видео.
        pass


# =========================================================
# ЗАПУСК БОТА
# =========================================================

if __name__ == "__main__":
    health_thread = threading.Thread(target=run_health, daemon=True)
    health_thread.start()
    print("Nova SI запущен.")
    bot.infinity_polling(skip_pending=True, timeout=30, long_polling_timeout=30)
